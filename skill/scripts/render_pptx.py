"""layout.json → 편집 가능한 .pptx.

    python render_pptx.py --layout layout.json --out map.pptx
    python render_pptx.py --layout a.json b.json --out deck.pptx     # 여러 장

지도는 이미지로 넣고, **카드·배지·지시선·마커·범례는 도형**으로 만든다.
파워포인트에서 문구를 고치고 핀을 옮길 수 있어야 하기 때문이다.

좌표는 layout.json 을 그대로 쓴다. 여기서 배치를 다시 계산하면 HTML 출력과 갈라진다.
python-pptx 가 필요하다 — `skill/.venv/bin/python` 으로 실행할 것.
"""
from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from common import design_tokens, log, read_json

EMU_IN = 914400


def I(v: float) -> Emu:
    return Emu(int(round(v * EMU_IN)))


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#").upper())


def _set_typefaces(run, latin: str, east_asian: str | bool | None = None,
                   complex_font: str | None = None) -> None:
    """python-pptx가 빠뜨리는 East Asian/complex 글꼴 정보를 명시한다.

    `font.name`만 설정하면 OOXML의 ``a:latin``만 생긴다. PowerPoint for Mac은
    이 경우 한글을 다른 글꼴로 대체해 샘플보다 크고 낮게 렌더한다.
    """
    run.font.name = latin
    rpr = run._r.get_or_add_rPr()
    ea_face = latin if east_asian is None else (
        None if east_asian is False else east_asian
    )
    for tag, face in (
        ("latin", latin),
        ("ea", ea_face),
        ("cs", complex_font),
    ):
        if not face:
            continue
        child = rpr.find(qn(f"a:{tag}"))
        if child is None:
            child = rpr.makeelement(qn(f"a:{tag}"))
            rpr.append(child)
        child.set("typeface", face)


def _set_run_metrics(run, spacing: int | None = None,
                     baseline: int | None = None) -> None:
    rpr = run._r.get_or_add_rPr()
    if spacing is not None:
        rpr.set("spc", str(spacing))
    if baseline is not None:
        rpr.set("baseline", str(baseline))
    rpr.set("dirty", "0")


# ─────────────────────────────── 지도 이미지 ───────────────────────────────

def map_image(layout_path: Path, tmp: Path, scale: int = 3) -> tuple[Path, dict]:
    """지도만 렌더해 PNG 로 굽고, 실제 그림 영역만 잘라낸다."""
    import export
    import render_html

    L = read_json(layout_path)
    tmp.mkdir(parents=True, exist_ok=True)
    html = tmp / "map_only.html"
    html.write_text(render_html.render(L, layers="map"), encoding="utf-8")
    png = export.to_png(html, tmp / "map_only.png", scale=scale)

    from PIL import Image
    c = L["map"].get("content") or L["map"]["frame"]
    px = 96 * scale
    im = Image.open(png)
    box = (max(int(c["x"] * px), 0), max(int(c["y"] * px), 0),
           min(int((c["x"] + c["w"]) * px), im.width),
           min(int((c["y"] + c["h"]) * px), im.height))
    out = tmp / "map_crop.png"
    im.crop(box).save(out)
    return out, c


# ─────────────────────────────── 도형 ───────────────────────────────

def textbox(slide, x, y, w, h, text, size, color, *, bold=False, align=PP_ALIGN.LEFT,
            font="맑은 고딕", anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing != 1.0:
            p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        _set_typefaces(r, font)
        r.font.color.rgb = rgb(color)
    return tb


def template_runs_textbox(slide, spec: dict, color: str,
                          font: str = "맑은 고딕"):
    """샘플2의 한 문단 자동맞춤 리치텍스트 상자를 그대로 만든다."""
    tb = slide.shapes.add_textbox(
        I(spec["x"]), I(spec["y"]), I(spec["w"]), I(spec["h"])
    )
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    # 샘플 OOXML은 anchor/alignment를 지정하지 않는다. add_textbox의 기본 anchor
    # 속성을 제거해 PowerPoint가 같은 글꼴 기준선으로 배치하게 한다.
    body_pr = tf._txBody.bodyPr
    body_pr.set("vert", "horz")
    body_pr.set("wrap", "square")
    body_pr.set("rtlCol", "0")
    body_pr.attrib.pop("anchor", None)

    p = tf.paragraphs[0]
    p_pr = p._p.get_or_add_pPr()
    p_pr.attrib.pop("algn", None)
    p_pr.set("marL", str(spec.get("paragraph_margin", 8145)))

    for run_spec in spec["runs"]:
        r = p.add_run()
        r.text = run_spec["text"]
        r.font.size = Pt(run_spec["size"])
        r.font.bold = True
        r.font.color.rgb = rgb(color)
        _set_typefaces(
            r,
            font,
            east_asian=run_spec.get("east_asian"),
            complex_font=run_spec.get("complex_font"),
        )
        _set_run_metrics(
            r,
            spacing=run_spec.get("spacing"),
            baseline=run_spec.get("baseline"),
        )
    return tb


def template_roman_textbox(slide, p: dict, roman: str, color: str):
    spec = {
        "x": p["label_x"], "y": p["label_y"],
        "w": p["label_w"], "h": p["label_h"],
        "paragraph_margin": p.get("paragraph_margin", 8145),
        "runs": [
            {
                "text": ch,
                "size": p.get("label_size", 8.98),
                "spacing": p.get("label_spacing", 58),
                "complex_font": "Book Antiqua",
                "east_asian": False,
            }
            for ch in roman
        ],
    }
    return template_runs_textbox(slide, spec, color, font="Georgia")


def flatten(shape):
    """도형에 붙는 테마 스타일(<p:style>)을 떼어 그림자·광택을 없앤다.

    `shadow.inherit = False` 는 빈 <a:effectLst/> 만 넣는데, 뷰어에 따라
    테마의 effectRef 를 그대로 살려 그림자가 남는다. 원본 샘플은 평면이다.
    """
    el = shape._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if el is not None:
        shape._element.remove(el)
    return shape


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = flatten(slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h)))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def draw_slide(prs, L: dict, tok: dict, tmp: Path, layout_path: Path) -> None:
    col, cd = tok["color"], tok["card"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1) 지도 — 이미지
    img, c = map_image(layout_path, tmp)
    slide.shapes.add_picture(str(img), I(c["x"]), I(c["y"]), I(c["w"]), I(c["h"]))

    # 2) 지시선 — 꺾은선은 직선 커넥터 여러 개로 나눠 그린다.
    #    python-pptx 의 ELBOW 커넥터는 꺾이는 지점을 지정할 수 없어 쓸 수 없다.
    for ld in L["leaders"]:
        pts = ld.get("points") or [ld["from"], ld["to"]]
        for p, q in zip(pts, pts[1:]):
            cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                            I(p[0]), I(p[1]), I(q[0]), I(q[1]))
            cn.line.color.rgb = rgb(col["leader"])
            cn.line.width = Pt(tok["leader"]["w_pt"])

    # 3) 마커
    mk = tok["marker"]
    for m in L["markers"]:
        if m.get("marker_kind") == "multi":
            d = mk["area"]["d"]
            rect(slide, m["x"] - d / 2, m["y"] - d / 2, d, d, None,
                 col["area_ring"], mk["area"]["line_w"], MSO_SHAPE.OVAL)
        else:
            d = mk["point"]["d"]
            rect(slide, m["x"] - d / 2, m["y"] - d / 2, d, d, col["pin"], None,
                 shape=MSO_SHAPE.OVAL)

    # 4) 카드
    bd = cd["badge"]
    for card in L["cards"]:
        x, y = card["x"], card["y"]
        body_dy = card.get("body_dy", cd["name"]["dy"])
        if card.get("show_place", True):
            textbox(slide, x, y, cd["name"]["w"], cd["place"]["h"] * 1.4,
                    card["place"], card.get("place_font", cd["place"]["size"]),
                    col["heading"], bold=True)
        bx = x + bd["dx"]
        for k in card["badges"]:
            by = y + body_dy + bd["dy"] - cd["name"]["dy"]
            rect(slide, bx, by, bd["w"], bd["h"], col["badge_bg"])
            textbox(slide, bx, by, bd["w"], bd["h"], _sym(k), bd["size"],
                    col["badge_fg"], bold=True, align=PP_ALIGN.CENTER,
                    anchor=MSO_ANCHOR.MIDDLE)
            bx += bd["w"] + bd["gap"]
        # 줄바꿈은 layout.py 가 이미 정했다. 여기서 다시 흘리면 HTML 과 달라진다.
        textbox(slide, x + cd["name"]["dx"], y + body_dy - 0.015,
                cd["name"]["w"], card["h"] - body_dy + 0.02, card["lines"], card["font"],
                col["body"], spacing=0.92)
        # 다중 대상지 사업의 설명 — 그 카드 안에 붙는다 (원본과 동일)
        if card.get("note"):
            textbox(slide, x + cd["name"]["dx"],
                    y + body_dy - 0.015 + len(card["lines"]) * card["line_h"],
                    cd["name"]["w"], card["line_h"] * 1.3, card["note"], card["font"],
                    col["area_ring"], bold=True)

    draw_title(slide, L, tok)
    draw_region_tab(slide, L, tok)
    draw_legend(slide, L, tok)
    


_SYMS = {"E": "E", "H": "H", "G": "G", "A": "A", "T": "…"}


def _sym(k: str) -> str:
    return _SYMS.get(k, k)


_PIN_OUTER_PATH = [
    (176911, 0), (129878, 6319), (87616, 24154), (51812, 51817),
    (24151, 87622), (6318, 129882), (0, 176910), (1180, 197439),
    (4632, 217276), (10222, 236291), (17818, 254355), (156095, 552310),
    (159867, 560044), (167754, 565378), (176911, 565378), (335572, 255181),
    (349069, 217736), (353822, 176910), (347502, 129882), (329667, 87622),
    (302004, 51817), (266199, 24154), (223939, 6319), (176911, 0),
]

_PIN_INNER_PATH = [
    (145148, 0), (99270, 7401), (59425, 28009), (28005, 59434),
    (7399, 99281), (0, 145160), (7399, 191040), (28005, 230887),
    (59425, 262312), (99270, 282920), (145148, 290321), (191032, 282920),
    (230880, 262312), (262303, 230887), (282909, 191040), (290309, 145160),
    (282909, 99281), (262303, 59434), (230880, 28009), (191032, 7401),
    (145148, 0),
]


def _source_freeform(slide, x: float, y: float, w: float, h: float,
                     points: list[tuple[int, int]], path_w: int, path_h: int,
                     fill: str):
    """샘플2 자유형의 꼭짓점을 그대로 재현한다."""
    mapped = [
        (x + px / path_w * w, y + py / path_h * h)
        for px, py in points
    ]
    ff = slide.shapes.build_freeform(I(mapped[0][0]), I(mapped[0][1]))
    ff.add_line_segments([(I(px), I(py)) for px, py in mapped[1:]], close=True)
    s = flatten(ff.convert_to_shape())
    # 원본 path 좌표의 마지막 점이 path w/h보다 약간 안쪽이므로 xfrm은 실측
    # 바운드로 다시 고정한다.
    s.left, s.top, s.width, s.height = I(x), I(y), I(w), I(h)
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def map_pin(slide, p: dict, fill: str):
    """샘플2의 보라 물방울 자유형을 동일한 꼭짓점으로 그린다."""
    return _source_freeform(
        slide, p["x"], p["y"], p["w"], p["h"],
        _PIN_OUTER_PATH, 354330, 565785, fill,
    )


def map_pin_inner(slide, p: dict):
    """샘플2의 흰 내부 타원도 프리셋 원이 아닌 원본 자유형으로 그린다."""
    return _source_freeform(
        slide,
        p["inner_x"], p["inner_y"], p["inner_w"], p["inner_h"],
        _PIN_INNER_PATH, 290830, 290830, "FFFFFF",
    )


def outline_box(slide, x: float, y: float, w: float, h: float,
                color: str, line_w: float):
    """인간 작성본과 같은 자유형 사각 테두리."""
    ff = slide.shapes.build_freeform(I(x), I(y))
    ff.add_line_segments([
        (I(x + w), I(y)),
        (I(x + w), I(y + h)),
        (I(x), I(y + h)),
    ], close=True)
    shape = flatten(ff.convert_to_shape())
    shape.fill.background()
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    return shape


def draw_title(slide, L: dict, tok: dict) -> None:
    t, col = tok["title"], tok["color"]
    spec = L.get("title") or {}
    p, b = t["pin"], spec.get("box") or t["box"]

    # 원본 z-order: 제목 텍스트 → 박스 → 물방울 → 흰 타원 → 로마숫자.
    # 박스를 먼저 그리면 긴 `ㅣ` 획이 아래선을 검게 덮어 사용자 눈에 겹침으로 보인다.
    if spec and spec.get("runs"):
        template_runs_textbox(
            slide, spec, col.get("title", col["heading"]), font="맑은 고딕"
        )
    else:  # 예전 layout.json 호환
        text = L["country"]["ko"] if L["lang"] == "ko" else L["country"]["en"]
        textbox(slide, b["x"] + 0.69, b["y"], b["w"] - 0.7, b["h"],
                text, 14.11, col["heading"], bold=True, anchor=MSO_ANCHOR.MIDDLE)

    outline_box(slide, b["x"], b["y"], b["w"], b["h"],
                col["region_tab"], b.get("line_w", 1.2))
    # 핀 = 보라 물방울 + 흰 타원 + 별도 로마숫자 상자. 세 도형의 실측 좌표가
    # 서로 다르므로 원 중심으로 재계산하지 않는다.
    map_pin(slide, p, col["region_tab"])
    map_pin_inner(slide, p)
    template_roman_textbox(
        slide, p, _roman(L.get("index", 1)), col["region_tab"]
    )


def _roman(n: int) -> str:
    out = ""
    for v, s in [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]:
        while n >= v:
            out += s
            n -= v
    return out or "I"


def draw_region_tab(slide, L: dict, tok: dict) -> None:
    if not L.get("region"):
        return
    r, col = tok["region_tab"], tok["color"]
    rect(slide, r.get("rail_x", r["x"]), r.get("rail_y", r["y"]),
         r.get("rail_w", r["w"]), r.get("rail_h", r["h"]), col["region_tab_bg"])
    tb = textbox(slide, r["x"] - (r["h"] - r["w"]) / 2, r["y"] + (r["h"] - r["w"]) / 2,
                 r["h"], r["w"], L["region"], r["size"], col["region_tab"],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tb.rotation = 90


def draw_legend(slide, L: dict, tok: dict) -> None:
    """범례 — 좌표·글자 크기는 layout._legend 가 정한다. 여기서 재계산하지 말 것.

    예전엔 HTML 만 넘침 축소를 했고 PPTX 는 빠뜨려, 영문 범례가 슬라이드 밖으로
    흘러나갔다. 두 렌더러가 같은 값을 쓰도록 배치를 layout 으로 올렸다.
    """
    lg, col = L["legend"], tok["color"]
    for it in lg["items"]:
        rect(slide, it["x"], it["y"], lg["swatch"], lg["swatch"], col["badge_bg"])
        textbox(slide, it["x"], it["y"], lg["swatch"], lg["swatch"], it["symbol"],
                lg["size"], col["badge_fg"], bold=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        cx, h = it["x"] + lg["swatch"] / 2, it["label_h"]
        if it["rotate"]:                       # 영문: 가로 텍스트박스를 90° 회전
            tb = textbox(slide, cx - h / 2, it["y"] + lg["swatch"] + h / 2 - 0.06,
                         h, 0.12, it["label"], it["size"], col["body"],
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            tb.rotation = 90
        else:                                  # 한글: 글자를 세로로 쌓는다
            textbox(slide, it["x"] - 0.03, it["y"] + lg["swatch"] + 0.02,
                    lg["swatch"] + 0.06, h, list(it["label"]), it["size"],
                    col["body"], align=PP_ALIGN.CENTER, spacing=0.9)
    pn = lg["page_num"]
    textbox(slide, pn["x"], pn["y"] - 0.14, 0.3, 0.2, f'{L.get("index", 1):02d}',
            pn["size"], col["body"])


def draw_notes(slide, L: dict, tok: dict) -> None:
    if not L.get("notes"):
        return
    lay = tok["layout"][L["mode"]]
    if L["mode"] == "A":
        x, y = lay["bottom_row"]["x0"], lay["bottom_row"]["y"] - 0.24
    else:
        x, y = lay["right_col"]["x"], lay["right_col"]["y1"] + 0.05
    textbox(slide, x, y, 2.0, 0.18, " ".join(L["notes"]), 7.0,
            tok["color"]["area_ring"], bold=True)


# ─────────────────────────────── CLI ───────────────────────────────

def build(layout_paths: list, out: Path) -> Path:
    tok = design_tokens()
    prs = Presentation()
    prs.slide_width = I(tok["canvas"]["w_in"])
    prs.slide_height = I(tok["canvas"]["h_in"])
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        for i, lp in enumerate(layout_paths, 1):
            L = read_json(lp)
            log(f'  [{i}/{len(layout_paths)}] {L["country"]["ko"]} ({L["lang"]})')
            draw_slide(prs, L, tok, tmp / f"s{i}", Path(lp))
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    log(f"  → {out} ({out.stat().st_size//1024}KB, 슬라이드 {len(layout_paths)}장)")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description="layout.json → 편집 가능한 PPTX")
    ap.add_argument("--layout", required=True, nargs="+", help="layout.json (여러 개 가능)")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    build([Path(p) for p in a.layout], Path(a.out))
    return 0


if __name__ == "__main__":
    sys.exit(main())

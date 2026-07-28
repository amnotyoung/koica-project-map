from __future__ import annotations

import sys
import unittest
from pathlib import Path


SCRIPTS = Path(__file__).resolve().parents[1] / "skill" / "scripts"
sys.path.insert(0, str(SCRIPTS))

from common import design_tokens
import layout

try:
    from pptx import Presentation
    from pptx.oxml.ns import qn
    import render_pptx
except ImportError:
    Presentation = None


@unittest.skipIf(Presentation is None, "python-pptx is installed in skill/.venv")
class TitlePptxTests(unittest.TestCase):
    def test_korean_title_is_one_inline_textbox_with_east_asian_font(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tok = design_tokens()
        spec = layout._title_layout("피지", "Fiji", "ko", tok)
        render_pptx.draw_title(
            slide,
            {
                "index": 2,
                "lang": "ko",
                "country": {"ko": "피지", "en": "Fiji"},
                "title": spec,
            },
            tok,
        )

        title_shapes = [shape for shape in slide.shapes if shape.text == "피지 Fiji"]
        self.assertEqual(len(title_shapes), 1)
        self.assertFalse(any(shape.text == "피지" for shape in slide.shapes))
        self.assertFalse(any(shape.text == "Fiji" for shape in slide.shapes))

        title = title_shapes[0]
        self.assertEqual(len(title.text_frame.paragraphs), 1)
        self.assertNotIn("\n", title.text)
        shapes = list(slide.shapes)
        border = next(
            shape for shape in shapes
            if abs(shape.left / 914400 - 2.27824) < 0.0001
            and abs(shape.top / 914400 - 0.46716) < 0.0001
        )
        self.assertLess(shapes.index(title), shapes.index(border))
        self.assertAlmostEqual(title.left / 914400, 2.96902, places=5)
        self.assertAlmostEqual(title.top / 914400, 0.52466, places=5)
        self.assertAlmostEqual(title.width / 914400, 0.68294, places=5)
        self.assertAlmostEqual(title.height / 914400, 0.25910, places=5)
        self.assertTrue(title.text_frame.word_wrap)

        runs = title.text_frame.paragraphs[0].runs
        self.assertEqual([run.text for run in runs], ["피지 ", "Fiji"])
        first_rpr = runs[0]._r.get_or_add_rPr()
        self.assertEqual(first_rpr.get("spc"), "-96")
        self.assertEqual(first_rpr.get("baseline"), "-5787")
        self.assertEqual(
            first_rpr.find(qn("a:ea")).get("typeface"),
            "맑은 고딕",
        )

        roman = next(shape for shape in slide.shapes if shape.text == "II")
        self.assertAlmostEqual(roman.left / 914400, 2.19732, places=5)
        self.assertAlmostEqual(roman.top / 914400, 0.50307, places=5)
        roman_rpr = roman.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
        self.assertIsNone(roman_rpr.find(qn("a:ea")))


if __name__ == "__main__":
    unittest.main()
